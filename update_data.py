import base64
import io
import json
import requests
from openpyxl import load_workbook

ONEDRIVE_URL = "https://1drv.ms/x/c/0a8fa620daf4b126/IQBGkzekdy3FR5DLok9EPZgJAQKIa0mG4ldlZVZ4IzE6SZU"
OUT = "data.json"

def download_xlsx():
    headers = {"User-Agent": "Mozilla/5.0"}
    urls = [
        ONEDRIVE_URL + "?download=1",
        ONEDRIVE_URL,
    ]
    last_error = None

    for url in urls:
        try:
            r = requests.get(url, headers=headers, timeout=60, allow_redirects=True)
            r.raise_for_status()

            if r.content[:2] == b"PK":
                return r.content

            # OneDrive may return an HTML page. Look for a direct download URL.
            text = r.content.decode("utf-8", errors="ignore")
            markers = ['downloadUrl":"', 'downloadUrl": "', 'href="']
            for marker in markers:
                pos = text.find(marker)
                if pos < 0:
                    continue
                start = pos + len(marker)
                end = text.find('"', start)
                if end < 0:
                    continue

                direct = text[start:end]
                direct = direct.replace("\\u0026", "&").replace("\\/", "/")
                if direct.startswith("http"):
                    rr = requests.get(
                        direct,
                        headers=headers,
                        timeout=60,
                        allow_redirects=True,
                    )
                    rr.raise_for_status()
                    if rr.content[:2] == b"PK":
                        return rr.content

            last_error = RuntimeError(
                f"OneDrive did not return an XLSX file (HTTP {r.status_code})."
            )
        except Exception as exc:
            last_error = exc

    raise RuntimeError(
        "Could not download the OneDrive Excel file. "
        f"Last error: {last_error}"
    )

def clean(value):
    return 0 if value is None or value == "" else value

def main():
    xlsx = download_xlsx()

    wb = load_workbook(
        io.BytesIO(xlsx),
        data_only=True,
        read_only=True,
    )

    if "FINAL" not in wb.sheetnames:
        raise RuntimeError(
            "FINAL sheet not found. Available sheets: "
            + ", ".join(wb.sheetnames)
        )

    ws = wb["FINAL"]
    records = []

    for row in range(5, ws.max_row + 1):
        badge = ws.cell(row, 2).value
        if badge is None or str(badge).strip() == "":
            continue

        name = ws.cell(row, 3).value
        parent = ws.cell(row, 4).value

        records.append({
            "badge": str(badge).strip().upper(),
            "name": "" if name is None else str(name).strip(),
            "parent": "" if parent is None else str(parent).strip(),
            "day": clean(ws.cell(row, 8).value),
            "night": clean(ws.cell(row, 9).value),
            "saturday": clean(ws.cell(row, 10).value),
            "sunday": clean(ws.cell(row, 11).value),
            "lipai": clean(ws.cell(row, 12).value),
            "beas": clean(ws.cell(row, 13).value),
        })

    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(records, f, ensure_ascii=False, indent=2)

    print(f"Wrote {len(records)} records to {OUT}")

if __name__ == "__main__":
    main()







